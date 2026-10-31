"""网页版从头到尾点一遍（真浏览器 + 假谷歌服务器），出的表跟电脑版一格一格对。

跑：python web/test/e2e.py [截图放哪]
要：playwright（chromium）、openpyxl、Pillow、xlwt；可选 pillow-heif（测 HEIC）
"""
import functools
import http.server
import io
import json
import os
import re
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
import zipfile
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, 'tests'))
from bangong import engine, tables  # noqa: E402
import mock_server  # noqa: E402
from test_kb import company_book  # noqa: E402
from test_formats import to_xls  # noqa: E402

K1, K2 = 'AIzaKEY1_mock', 'AIzaKEY2_mock'
RELAY_DIR = os.environ.get('RELAY_DIR') or os.path.join(os.path.dirname(ROOT), 'luruseng-relay')   # 中转站代码（QIANSHOULAB/luruseng-relay）
SHOTS = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix='luruseng-shots-')
os.makedirs(SHOTS, exist_ok=True)
WORK = tempfile.mkdtemp(prefix='luruseng-e2e-')
CORS = {'access-control-allow-origin': '*', 'access-control-allow-methods': 'POST, OPTIONS',
        'access-control-allow-headers': 'content-type, x-goog-api-key'}
passed = []


def step(fn):
    @functools.wraps(fn)
    def run(*a, **kw):
        t = time.time()
        fn(*a, **kw)
        passed.append(fn.__name__)
        print(f'  ✓ {fn.__name__}  ({time.time() - t:.1f}s)', flush=True)
    return run


def make_images(folder, n, start=0):
    os.makedirs(folder, exist_ok=True)
    out = []
    for i in range(start, start + n):
        p = os.path.join(folder, f'img_{i:02d}.jpg')
        Image.new('RGB', (800 + i, 600), (i * 7 % 255, 100, 150)).save(p)
        os.utime(p, (1_700_000_000 + i, 1_700_000_000 + i))
        out.append(p)
    return out


def serve(root):
    class Quiet(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *a):
            pass
    Quiet.extensions_map = {**http.server.SimpleHTTPRequestHandler.extensions_map, '.js': 'text/javascript',
                            '.mjs': 'text/javascript', '.webmanifest': 'application/manifest+json'}
    srv = http.server.ThreadingHTTPServer(('127.0.0.1', 0), functools.partial(Quiet, directory=root))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def start_relay():
    """本地起一个真的中转站（wrangler dev，不用 Cloudflare 账号）。没有中转站代码就返回 (None, None)，手机传图那项跳过。"""
    if not os.path.exists(os.path.join(RELAY_DIR, 'wrangler.jsonc')) or not os.path.isdir(os.path.join(RELAY_DIR, 'node_modules')):
        assert not os.environ.get('E2E_NEED_RELAY'), f'没有中转站代码：{RELAY_DIR}（先 npm ci）'
        return None, None
    with socket.socket() as so:
        so.bind(('127.0.0.1', 0))
        port = so.getsockname()[1]
    proc = subprocess.Popen(['npx', 'wrangler', 'dev', '--ip', '127.0.0.1', '--port', str(port), '--var', 'ALLOW_ORIGINS:*',
                             '--persist-to', tempfile.mkdtemp(prefix='luruseng-relay-')], cwd=RELAY_DIR,
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True,
                            env={**os.environ, 'WRANGLER_SEND_METRICS': 'false', 'CI': '1'})
    url = f'http://127.0.0.1:{port}'
    for _ in range(300):
        try:
            with urllib.request.urlopen(url + '/ping', timeout=2) as r:
                if r.status == 200:
                    return proc, url
        except Exception:
            time.sleep(0.3)
    os.killpg(proc.pid, signal.SIGTERM)
    raise AssertionError('中转站（wrangler dev）没起来')


def python_book(images, spec_name=None):
    """电脑版拿同样的照片、同样的假服务器出一本表：拿来跟网页版的一格一格对。"""
    d = tempfile.mkdtemp(prefix='luruseng-py-')
    for p in images:
        shutil.copy2(p, d)
    kw = {}
    if spec_name:
        cat = os.path.join(d, '_cat')
        kw['spec'] = tables.TableSpec(spec_name, os.path.join(cat, 'db.sqlite'), os.path.join(cat, 'feed'))
    res = engine.Engine(d, [K1], log=lambda m: None, **kw).run()
    assert res.ok == len(images), res
    return res.xlsx


def cells(path):
    ws = load_workbook(path).worksheets[0]
    out = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in row:
            fill = c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type == 'solid' else None
            font = (c.font.name, float(c.font.sz or 11), bool(c.font.b), c.font.color.rgb if c.font.color is not None and c.font.color.type == 'rgb' else None)
            v = c.value
            out[c.coordinate] = (v, fill, font, c.alignment.horizontal, bool(c.alignment.wrap_text), c.number_format if v is not None else None,
                                 bool(c.border.left.style))
    widths = {}
    for k, v in ws.column_dimensions.items():            # <col min max> 一段算好几列
        for i in range(v.min or column_index_from_string(k), (v.max or column_index_from_string(k)) + 1):
            if v.width:
                widths[get_column_letter(i)] = round(v.width, 2)
    meta = dict(title=ws.title, merged=sorted(str(m) for m in ws.merged_cells.ranges), freeze=ws.freeze_panes,
                filter=ws.auto_filter.ref, orient=ws.page_setup.orientation, fitw=ws.page_setup.fitToWidth,
                fith=ws.page_setup.fitToHeight, fit=ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None,
                titles=ws.print_title_rows, height1=ws.row_dimensions[1].height)
    return out, widths, meta


def same_book(js_path, py_path, skip=()):
    a, wa, ma = cells(js_path)
    b, wb, mb = cells(py_path)
    diffs = []
    for k in sorted(set(a) | set(b)):
        x, y = a.get(k), b.get(k)
        if x is None or y is None:            # 一边没有这个格子：只要另一边是空白就算一样
            z = x or y
            if z[0] is not None or z[1] is not None:
                diffs.append((k, x, y))
            continue
        if x[0] != y[0] and isinstance(x[0], str) and isinstance(y[0], str) and any(s in x[0] for s in skip):
            continue
        if x != y:
            diffs.append((k, x, y))
    for k in set(wa) | set(wb):
        if abs(wa.get(k, 0) - wb.get(k, 0)) > 0.01:
            diffs.append(('width ' + k, wa.get(k), wb.get(k)))
    for k in ma:
        if ma[k] != mb[k] and not (k == 'titles' and str(ma[k]).replace('$', '') == str(mb[k]).replace('$', '')):
            diffs.append((k, ma[k], mb[k]))
    assert not diffs, '网页版和电脑版的表不一样：\n' + '\n'.join(map(str, diffs[:30]))


def main():
    site = serve(ROOT)
    base = f'http://127.0.0.1:{site.server_address[1]}/web/'
    mock = mock_server.start()
    mock_url = f'http://127.0.0.1:{mock.server_address[1]}'
    engine.API_BASE = mock_url + '/v1beta'
    engine.RETRY_WAITS = [0.01, 0.01, 0.01]
    mock_server.STATE.update(delay=0)
    errors = []
    relay_proc, relay_url = start_relay()

    def gemini(route, request):
        if request.method == 'OPTIONS':
            route.fulfill(status=204, headers=CORS)
            return
        u = request.url.split('googleapis.com', 1)[1]
        req = urllib.request.Request(mock_url + u, data=request.post_data_buffer, method='POST',
                                     headers={'Content-Type': 'application/json'})
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                status, body = r.status, r.read()
        except urllib.error.HTTPError as e:
            status, body = e.code, e.read()
        try:
            route.fulfill(status=status, headers={**CORS, 'content-type': 'application/json'}, body=body)
        except Exception:                     # 页面已经被关了（测「页面被杀」的时候）
            pass

    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        ctx = browser.new_context(viewport={'width': 1280, 'height': 860}, accept_downloads=True, locale='zh-CN')
        ctx.grant_permissions(['clipboard-read', 'clipboard-write'], origin=base.rsplit('/web/', 1)[0])
        page = ctx.new_page()
        page.on('pageerror', lambda e: errors.append(f'pageerror: {e}'))
        page.on('console', lambda m: errors.append(f'console.{m.type}: {m.text}')
                if m.type == 'error' and 'Failed to load resource' not in m.text else None)
        def offline(pg):                     # 除了（假的）谷歌，别的外网一概不许碰：网页版不靠 CDN
            pg.route('https://**', lambda route, req: errors.append(f'外网请求：{req.url}') or route.abort())
            pg.route('https://generativelanguage.googleapis.com/**', gemini)
        offline(page)

        def logs():
            return page.eval_on_selector_all('#log p', 'ps => ps.map(p => p.textContent)')

        def wait_log(text, timeout=60000):
            page.wait_for_function('t => [...document.querySelectorAll("#log p")].some(p => p.textContent.includes(t))',
                                   arg=text, timeout=timeout)

        def last_logs(n=8):
            return '\n'.join(logs()[-n:])

        def add_photos(paths):
            page.set_input_files('#files', paths)

        def run_and_wait(paths, until='成功', timeout=60000):
            before = len(logs())
            add_photos(paths)
            page.click('#start')
            page.wait_for_function('([t, n]) => [...document.querySelectorAll("#log p")].slice(n).some(p => p.textContent.includes(t))',
                                   arg=[until, before], timeout=timeout)
            page.wait_for_function('() => !document.querySelector("#start").disabled', timeout=timeout)
            return logs()[before:]

        def download(selector):
            with page.expect_download() as d:
                page.click(selector)
            path = os.path.join(WORK, d.value.suggested_filename)
            d.value.save_as(path)
            return path

        def wait_viewer():                    # 核对窗口开了、这批也读出来了
            page.wait_for_selector('#dlg-view[open]')
            page.wait_for_function('() => !!document.querySelector("#vpanel h3") || !document.querySelector("#vempty").hidden')

        def close_viewer():
            if page.is_visible('#dlg-view'):
                page.click('#vclose')

        @step
        def opens_clean():
            page.goto(base)
            page.wait_for_function('() => document.querySelectorAll("#log p").length >= 2')
            assert page.title().startswith('录入僧网页版')
            assert page.inner_text('#status').replace('\n', '') == '待识别 0 张', page.inner_text('#status')
            text = '\n'.join(logs())
            assert '第一次用' in text and '数据库还是空的' in text, text
            assert page.inner_text('#cathint') == '内置 · 专属优化'
            assert page.inner_text('#feed').strip() == '放入样板资料'
            page.screenshot(path=os.path.join(SHOTS, '1-打开.png'))

        @step
        def no_key_no_start():
            add_photos(make_images(os.path.join(WORK, 'none'), 1, start=90))
            page.click('#start')
            wait_log('✗ 先填钥匙')
            page.click('#clear')
            assert page.inner_text('#status').replace('\n', '') == '待识别 0 张'

        imgs = make_images(os.path.join(WORK, 'a'), 12)

        @step
        def twelve_photos_match_desktop_book():
            page.fill('#key1', K1)
            out = run_and_wait(imgs, '成功 12 张')
            text = '\n'.join(out)
            assert '按拍摄顺序编号 1–12' in text and '每张认两遍对答案' in text, text
            assert sum(1 for m in out if m.startswith('✓ ')) >= 12, text
            wait_viewer()                                        # 干完自动打开核对
            assert '第 1 张 / 共 12 张' in page.inner_text('#vinfo')
            panel = page.inner_text('#vpanel')
            assert 'No. 1' in panel and '800' in panel and 'VM-07' in panel, panel
            page.wait_for_function('() => { const i = document.querySelector("#vphoto img"); return i && i.naturalWidth > 0; }')
            page.screenshot(path=os.path.join(SHOTS, '2-核对.png'))
            page.keyboard.press('ArrowRight')                                    # 一打开就能用 ←→ 翻页
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("第 2 张")')
            page.keyboard.press('ArrowLeft')
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("第 1 张")')
            js = download('#vxlsx')
            assert os.path.basename(js).startswith('识别结果_') and js.endswith('.xlsx')
            same_book(js, python_book(imgs))
            assert len(mock_server.STATE['log']) >= 12 * 2 * 2      # 网页 24 次 + 电脑版 24 次

        @step
        def viewer_pages_rotates_copies_zips():
            page.keyboard.press('ArrowRight')
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("第 2 张")')
            assert '801' in page.inner_text('#vpanel')
            w0 = page.eval_on_selector('#vphoto img', 'i => [i.naturalWidth, i.naturalHeight]')
            page.click('#vcw')
            page.wait_for_function('w => { const i = document.querySelector("#vphoto img"); return i && i.naturalWidth === w[1] && i.naturalHeight === w[0]; }', arg=w0)
            page.fill('#vnum', '12')
            page.press('#vnum', 'Enter')
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("第 12 张")')
            page.fill('#vnum', '40')
            page.click('#vgo')
            assert '没有第 40 张（这批是 1–12）' in page.inner_text('#vinfo')
            page.click('#vcopy')
            page.wait_for_function('() => document.querySelector("#vinfo").textContent.includes("复制好了")')
            clip = page.evaluate('navigator.clipboard.readText()').split('\r\n')
            assert len(clip) == 12 and clip[0].split('\t')[:6] == ['2026-10-28', 'VM-07', 'TD-524613-5', '82664580', '800', 'T4'], clip[0]
            z = download('#vzip')
            names = sorted(zipfile.ZipFile(z).namelist())
            assert names[:12] == [f'{i:02d}.jpg' for i in range(1, 13)] and names[12].endswith('.xlsx'), names
            im = Image.open(io.BytesIO(zipfile.ZipFile(z).read('01.jpg'))).convert('RGB')
            reds = sum(1 for x in range(im.width) for y in range(0, im.height, 4)
                       if (lambda p: p[0] > 170 and p[1] < 60 and p[2] < 60)(im.getpixel((x, y))))
            assert im.size == (800, 600) and reds > 100, (im.size, reds)        # 印了红编号，没被放大
            close_viewer()

        @step
        def same_photos_again_are_skipped():
            out = run_and_wait(imgs, '收工')
            text = '\n'.join(out)
            assert text.count('以前干过') == 12 and '第 1 张' in text and '列表里没有新照片，收工。' in text, text

        @step
        def result_button_downloads_latest():
            path = download('#result')
            assert load_workbook(path).worksheets[0].max_row == 15

        @step
        def feed_database_xlsx_xls_csv_html():
            book = company_book(os.path.join(WORK, 'TOOL BROKEN RECORD.xlsx'))
            xls = to_xls(book, os.path.join(WORK, '老表.xls'))
            csv_ = os.path.join(WORK, '刀单.csv')
            with open(csv_, 'wb') as f:
                f.write('DATE,MACHINE,PART NUMBER,MO,OPR NO.,T NO.,DESCRIPTION,BTMP,TOTAL SCREENED,PASSED,FAILED,SHIFT,SETTER\n'
                        '28/10,VM-31,TD-1,82660000,150,T4,EM2.5,1,1,1,0,D,700/Ah Seng\n'.encode('gbk'))
            html = os.path.join(WORK, '系统导出.xls')
            with open(html, 'w', encoding='utf-8') as f:
                f.write('<html><body><table><tr><td>DATE</td><td>MACHINE</td><td>PART NUMBER</td><td>MO</td><td>T NO.</td>'
                        '<td>DESCRIPTION</td><td>SETTER</td></tr><tr><td>29/10</td><td>VM-32</td><td>TD-2</td><td>82660001</td>'
                        '<td>T5</td><td>EM3</td><td>701/Kumar</td></tr></table></body></html>')
            broken = os.path.join(WORK, '坏表.xls')
            with open(broken, 'wb') as f:
                f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'xx' * 50)
            pdf = os.path.join(WORK, 'scan.pdf')
            with open(pdf, 'wb') as f:
                f.write(b'%PDF-1.4')
            ours = [p for p in os.listdir(WORK) if p.startswith('识别结果_') and p.endswith('.xlsx')][0]
            before = len(logs())
            page.click('#feed')
            assert '一条都不吃' in page.inner_text('#feeddesc') and '公司' not in page.inner_text('#feeddesc')
            assert page.inner_text('#dlg-feed h2') == '放入样板资料'
            page.set_input_files('#feedfiles', [os.path.join(WORK, ours), book, xls, csv_, html, broken, pdf])
            page.wait_for_function('n => [...document.querySelectorAll("#log p")].slice(n).some(p => p.textContent.startsWith("数据库："))', arg=before)
            out = '\n'.join(logs()[before:])
            assert f'✓ 吃进 {ours}：读到 12 行记录，新增 12 行' in out, out
            assert '✓ 吃进 TOOL BROKEN RECORD.xlsx：读到 8 行记录，新增 8 行' in out, out
            assert '✓ 吃进 老表.xls：读到 8 行记录，新增 0 行' in out, out
            assert '✓ 吃进 刀单.csv：读到 1 行记录，新增 1 行' in out, out
            assert '✓ 吃进 系统导出.xls：读到 1 行记录，新增 1 行' in out, out
            assert '✗ 坏表.xls：这本表坏了' in out and '✗ scan.pdf：只认' in out, out
            assert '数据库：22 条记录' in out, out

        @step
        def database_fixes_confusable_letters():
            mock_server.STATE['modes'] = {K1: 'confusable'}
            try:
                run_and_wait(make_images(os.path.join(WORK, 'c'), 1, start=20), '成功 1 张')
            finally:
                mock_server.STATE['modes'] = {}
            wait_viewer()
            fixed = page.eval_on_selector_all('#vpanel dd.fixed', 'ds => ds.map(d => d.textContent)')
            assert 'VM-07' in fixed and 'TD-524613-5' in fixed, fixed
            assert 'MACHINE VM-O7→VM-07（按数据库改）' in page.inner_text('#vpanel')
            page.screenshot(path=os.path.join(SHOTS, '3-按数据库改.png'))
            close_viewer()

        def idb(store, key):
            return page.evaluate("""([s, k]) => new Promise((ok, bad) => {
                const r = indexedDB.open('luruseng-web');
                r.onsuccess = () => { const q = r.result.transaction(s).objectStore(s).get(k); q.onsuccess = () => ok(q.result); q.onerror = () => bad(q.error); };
                r.onerror = () => bad(r.error);
            })""", [store, key])

        def cell_of(f):
            return tuple(page.eval_on_selector(f'#vpanel dd[data-f="{f}"]', 'd => [d.className, d.textContent]'))

        def vinfo_has(text):
            page.wait_for_function('t => document.querySelector("#vinfo").textContent.includes(t)', arg=text)

        @step
        def edit_cells_in_viewer():
            mock_server.STATE['modes'] = {K1: 'chaos'}                     # 几遍认出的机台都不一样：红，底下给候选
            try:
                run_and_wait(make_images(os.path.join(WORK, 'e'), 1, start=110), '成功 1 张')
            finally:
                mock_server.STATE['modes'] = {}
            wait_viewer()
            cls, text = cell_of('machine')
            assert 'red' in cls, (cls, text)
            page.click('#vpanel dd[data-f="machine"]')
            page.wait_for_selector('#vpanel .cell-ed .alt')
            alts = page.eval_on_selector_all('#vpanel .cell-ed .alt', 'bs => bs.map(b => b.textContent)')
            assert 'VM-02' in alts and len(alts) >= 2, alts
            assert page.input_value('#vpanel .cell-ed input') == text
            page.screenshot(path=os.path.join(SHOTS, '4-核对改格子.png'))
            page.click('#vpanel dd[data-f="machine"]')                     # 再点一下：收起
            assert page.query_selector('#vpanel .cell-ed') is None
            page.click('#vpanel dd[data-f="part_number"]')
            page.press('#vpanel .cell-ed input', 'Escape')                  # Esc 只收起格子，核对窗口不关
            assert page.query_selector('#vpanel .cell-ed') is None and page.is_visible('#dlg-view')
            page.click('#vpanel dd[data-f="machine"]')
            page.click('#vpanel .cell-ed .alt >> text=VM-02')
            vinfo_has('MACHINE 改成 VM-02')
            assert cell_of('machine') == ('edit human', 'VM-02'), cell_of('machine')
            assert '✓ 人核过' in page.inner_text('#vpanel')
            assert '进了样板资料' in page.inner_text('#vinfo'), page.inner_text('#vinfo')
            kb = idb('kv', 'kb')
            mine = [r for r in kb['records'] if r['source'].startswith('核对 ')]
            assert len(mine) == 1 and mine[0]['machine'] == 'VM-02' and mine[0]['sheet'] == '核对', mine
            js = download('#vxlsx')
            ws = load_workbook(js).worksheets[0]
            assert ws['C4'].value == 'VM-02' and ws['C4'].fill.fgColor.rgb.endswith('E2EFDA'), (ws['C4'].value, ws['C4'].fill.fgColor.rgb)
            assert ws['C4'].font.color.rgb.endswith('375623') and ws['O4'].value == '✓ 人核过', ws['O4'].value
            assert not ws['D4'].fill.fgColor.rgb.endswith('E2EFDA')                # 没动的格子不染
            close_viewer()

            mock_server.STATE['modes'] = {K1: 'mismatch'}                  # 10 张里 9 合格 0 不合格：三格都红
            try:
                run_and_wait(make_images(os.path.join(WORK, 'e2'), 1, start=111), '成功 1 张')
            finally:
                mock_server.STATE['modes'] = {}
            wait_viewer()
            assert all('red' in cell_of(f)[0] for f in ('total_screened', 'passed', 'failed'))
            page.click('#vpanel dd[data-f="total_screened"]')
            page.click('#vpanel .cell-ed >> text=对的')
            vinfo_has('TOTAL_SCREENED 核过了'.replace('_', ' '))
            assert cell_of('total_screened') == ('edit human', '10') and 'red' in cell_of('failed')[0]
            assert 'PASSED 9 + FAILED 0 ≠ TOTAL 10' in page.inner_text('#vpanel')
            page.click('#vpanel dd[data-f="failed"]')
            page.fill('#vpanel .cell-ed input', '1o')
            page.press('#vpanel .cell-ed input', 'Enter')
            vinfo_has('还不对')
            assert 'red' in cell_of('failed')[0] and 'FAILED 不是数字：1o' in page.inner_text('#vpanel')
            page.click('#vpanel dd[data-f="failed"]')
            page.fill('#vpanel .cell-ed input', '1')
            page.press('#vpanel .cell-ed input', 'Enter')
            vinfo_has('FAILED 改成 1')
            assert [cell_of(f)[0] for f in ('total_screened', 'passed', 'failed')] == ['edit human', 'edit', 'edit human']
            assert '✓ 人核过' in page.inner_text('#vpanel') and '进了样板资料' in page.inner_text('#vinfo')
            run_id = page.input_value('#vruns')
            close_viewer()
            page.click('#check')                                           # 关了再开：改过的还在
            wait_viewer()
            assert page.input_value('#vruns') == run_id
            assert cell_of('failed') == ('edit human', '1'), cell_of('failed')
            kb = idb('kv', 'kb')
            assert len([r for r in kb['records'] if r['source'].startswith('核对 ')]) == 2
            close_viewer()

        @step
        def btmp_copies_total():
            mock_server.STATE['modes'] = {K1: 'nobtmp'}                    # 换刀调试单：单子上没有 BTMP
            try:
                run_and_wait(make_images(os.path.join(WORK, 'b'), 1, start=112), '成功 1 张')
            finally:
                mock_server.STATE['modes'] = {}
            wait_viewer()
            assert page.query_selector('#vpanel dd[data-f="btmp"]') is None
            text = page.inner_text('#vpanel')
            assert 'BTMP' not in text and 'TOTAL SCREENED' in text, text
            hs = page.evaluate("""() => Object.fromEntries([...document.querySelectorAll('#vpanel dt')]
                .map(d => [d.textContent, Math.round(d.getBoundingClientRect().height)]))""")
            assert hs['TOTAL SCREENED'] == hs['MACHINE'], hs                  # 一行，不折成两行
            ws = load_workbook(download('#vxlsx')).worksheets[0]
            assert ws['I2'].value == 'BTMP' and ws['I4'].value == 1 == ws['J4'].value   # 表里 BTMP 抄 TOTAL SCREENED，贴进公司表就齐了
            assert (ws['I4'].fill.fgColor.rgb or '').endswith('FFE699')
            page.click('#vpanel dd[data-f="total_screened"]')              # 核对里改 TOTAL：BTMP 跟着变
            page.fill('#vpanel .cell-ed input', '3')
            page.press('#vpanel .cell-ed input', 'Enter')
            vinfo_has('TOTAL_SCREENED 改成 3'.replace('_', ' '))
            assert page.query_selector('#vpanel dd[data-f="btmp"]') is None
            ws = load_workbook(download('#vxlsx')).worksheets[0]
            assert ws['I4'].value == 3 == ws['J4'].value, (ws['I4'].value, ws['J4'].value)
            close_viewer()

        @step
        def names_dialog():
            page.click('#names')
            page.fill('#namestxt', '639/Lee Ah Kow\n# 注释\n232/Tan Mei Ling')
            page.click('#namesok')
            wait_log('✓ 名单存好了：2 人')
            assert page.inner_text('#nhead') == '名单：2 人'

        cat_imgs = make_images(os.path.join(WORK, 'q'), 2, start=30)

        @step
        def own_category_matches_desktop_book():
            page.click('#newcat')
            page.fill('#newname', '刀单')
            page.click('#newok')
            assert page.inner_text('#newerr') == '已经有「刀单」了'
            page.fill('#newname', '质检单')
            page.click('#newok')
            wait_log('「质检单」还没放过表')
            assert page.input_value('#cat') == '质检单' and page.inner_text('#cathint') == '自建 · 还没放过表'
            assert page.is_hidden('#names')
            out = run_and_wait(cat_imgs, '成功 2 张')
            assert any('类别「质检单」（还没定列：认全部内容）' in m for m in out), out
            wait_viewer()
            assert 'M/C No.' in page.inner_text('#vpanel')
            js = download('#vxlsx')
            assert js.endswith('_质检单.xlsx'), js
            same_book(js, python_book(cat_imgs, '质检单'))
            close_viewer()
            before = len(logs())
            page.click('#feed')
            assert page.input_value('#feedcat') == '质检单'
            page.set_input_files('#feedfiles', [js])
            page.wait_for_function('n => [...document.querySelectorAll("#log p")].slice(n).some(p => p.textContent.startsWith("「质检单」数据库："))', arg=before)
            out = '\n'.join(logs()[before:])
            assert '✓ 放进去' in out and '新增 2 行' in out and '新列：Date、M/C No.、Qty、Remark' in out, out
            assert page.inner_text('#cathint') == '自建 · 4 列 · 2 条记录'
            out = run_and_wait(make_images(os.path.join(WORK, 'q2'), 1, start=40), '成功 1 张')
            assert any('类别「质检单」（4 列）' in m for m in out), out
            close_viewer()
            page.select_option('#cat', '刀单')
            wait_log('数据库：')

        @step
        def bad_key_stops_and_keeps_photos():
            mock_server.STATE['modes'] = {K1: 'invalid'}
            try:
                out = run_and_wait(make_images(os.path.join(WORK, 'bad'), 2, start=50), '停工')
            finally:
                mock_server.STATE['modes'] = {}
            assert any('■ 停工：钥匙1 · HTTP 400：API key not valid' in m for m in out), out
            assert page.inner_text('#status').replace('\n', '') == '待识别 2 张'     # 照片还在列表里
            out = run_and_wait([], '成功 2 张')                                   # 钥匙好了再点开始：接着干
            assert any('成功 2 张' in m for m in out)
            close_viewer()

        @step
        def stop_button_marks_unfinished():
            mock_server.STATE['delay'] = 0.25
            try:
                add_photos(make_images(os.path.join(WORK, 's'), 12, start=60))
                page.click('#start')
                page.wait_for_function('() => /^5\\/12/.test(document.querySelector("#status").textContent)', timeout=30000)
                page.click('#stop')
                wait_log('■ 已停止，没干完的下次接着干', timeout=60000)
                page.wait_for_function('() => !document.querySelector("#start").disabled')
            finally:
                mock_server.STATE['delay'] = 0
            left = page.inner_text('#droptxt')
            assert left.startswith('7 张：img_65.jpg'), left                     # 没干完的还在列表里
            page.click('#check')
            wait_viewer()
            page.fill('#vnum', '6')
            page.press('#vnum', 'Enter')
            page.wait_for_function('() => document.querySelector("#vpanel").textContent.includes("No. 6")')
            assert '未完成：已停止' in page.inner_text('#vpanel'), page.inner_text('#vpanel')
            close_viewer()
            page.click('#clear')

        @step
        def photos_of_all_kinds():
            d = os.path.join(WORK, 'kinds')
            os.makedirs(d, exist_ok=True)
            big = os.path.join(d, 'big.png')
            Image.new('RGBA', (4032, 3024), (0, 0, 0, 0)).save(big)
            rot = os.path.join(d, 'rot.jpg')
            im = Image.new('RGB', (900, 600), 'white')
            ex = im.getexif()
            ex[0x0112] = 6                                  # 横着拍的：看的时候要转 90°
            im.save(rot, exif=ex)
            paths = [big, rot]
            try:
                import pillow_heif
                pillow_heif.register_heif_opener()
                heic = os.path.join(d, 'iphone.heic')
                Image.new('RGB', (1000, 700), (200, 220, 240)).save(heic, format='HEIF')
                paths.append(heic)
            except ImportError:
                print('    （没装 pillow-heif，HEIC 不测）')
            out = run_and_wait(paths, f'成功 {len(paths)} 张', timeout=120000)
            wait_viewer()
            js = download('#vxlsx')
            ws = load_workbook(js).worksheets[0]
            widths = sorted(ws.cell(row=r, column=6).value for r in range(4, 4 + len(paths)))
            want = sorted([2048, 600] + ([1000] if len(paths) == 3 else []))
            assert widths == want, (widths, out)
            close_viewer()

        @step
        def title_on_the_right_turns_upright_number_in_top_right():
            d = os.path.join(WORK, 'side')
            os.makedirs(d, exist_ok=True)
            p = os.path.join(d, 'side.jpg')
            Image.new('RGB', (900, 600), 'white').save(p)
            mock_server.STATE['modes'] = {K1: 'side_right'}
            try:
                run_and_wait([p], '成功 1 张')
            finally:
                mock_server.STATE['modes'] = {}
            wait_viewer()
            page.wait_for_function('() => { const i = document.querySelector("#vphoto img"); return i && i.complete && i.naturalWidth > 0; }')
            fits = page.evaluate('''() => { const b = document.querySelector("#vphoto").getBoundingClientRect();
                const i = document.querySelector("#vphoto img"), r = i.getBoundingClientRect();
                const s = Math.min(r.width / i.naturalWidth, r.height / i.naturalHeight);   // object-fit: contain 实际画多大
                return [i.naturalWidth * s <= b.width + 1, i.naturalHeight * s <= b.height + 1, r.top >= b.top - 1, r.bottom <= b.bottom + 1]; }''')
            assert all(fits), fits                                               # 竖的照片整张放得下，不被裁掉
            z = download('#vzip')
            im = Image.open(io.BytesIO(zipfile.ZipFile(z).read('01.jpg'))).convert('RGB')
            assert im.size == (600, 900), im.size                                 # 标题在右边：转成竖的
            red = [(x, y) for x in range(0, im.width, 2) for y in range(0, im.height, 2)
                   if (lambda q: q[0] > 170 and q[1] < 70 and q[2] < 70)(im.getpixel((x, y)))]
            corner = [1 for x, y in red if x > im.width * 0.6 and y < im.height * 0.3]
            assert red and len(corner) / len(red) > 0.95, (len(red), len(corner))   # 编号印在右上角，字是正的
            close_viewer()

        @step
        def settings_backup_restore_and_forget_keys():
            page.click('#gear')
            page.wait_for_selector('#dlg-set[open]')
            assert '录入僧网页版 v' in page.inner_text('#setver') and '千手院方丈释手' in page.inner_text('#setver')
            assert '数据库 24 条' in page.inner_text('#usage'), page.inner_text('#usage')   # 22 条放进去的 + 核对里核完的 2 条
            page.screenshot(path=os.path.join(SHOTS, '4-设置.png'))
            assert page.get_attribute('#hub', 'href') == 'https://qianshoulab.github.io/'
            page.eval_on_selector('#mail', 'a => a.addEventListener("click", e => e.preventDefault())')   # 测试里不真去开邮件程序
            page.click('#mail')                                                   # 方丈邮箱：先复制邮箱，底下一行写出来
            page.wait_for_function('() => document.querySelector("#setnote").textContent.includes("已复制")')
            assert page.get_attribute('#mail', 'href') is None                    # 只复制、亮出邮箱，不开邮件程序
            assert page.text_content('#setnote').startswith(page.get_attribute('#mail', 'title'))
            assert page.evaluate('navigator.clipboard.readText()') == page.get_attribute('#mail', 'title')
            assert page.get_attribute('#mail', 'title').startswith('oilsongchengmin@')
            path = download('#backup')
            data = json.load(open(path, encoding='utf-8'))
            assert data['app'] == 'luruseng-web' and len(data['kb']['records']) == 24 and '质检单' in data['cats']
            assert K1 not in json.dumps(data)                                     # 钥匙不进备份
            page.set_input_files('#restorefile', path)
            wait_log('✓ 恢复好了（合并进来，原来的一条没删）：数据库新增 0 条')
            page.click('#gear')
            page.wait_for_selector('#dlg-set[open]')
            page.click('#setrows label:nth-child(2) .switch')                    # 关掉「记住钥匙」
            page.click('#dlg-set [data-close]')
            page.reload()
            page.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
            assert page.input_value('#key1') == ''
            page.fill('#key1', K1)
            page.click('#gear')
            page.wait_for_selector('#dlg-set[open]')
            page.click('#setrows label:nth-child(2) .switch')
            page.click('#dlg-set [data-close]')
            page.reload()
            page.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
            assert page.input_value('#key1') == K1

        @step
        def delete_a_batch():
            page.click('#check')
            wait_viewer()
            n0 = page.eval_on_selector('#vruns', 's => s.options.length')
            page.click('#vdel')
            page.click('#dlg-ask [data-close]')
            assert page.eval_on_selector('#vruns', 's => s.options.length') == n0
            page.click('#vdel')
            page.click('#askok')
            page.wait_for_function('n => document.querySelector("#vruns").options.length === n - 1', arg=n0)
            close_viewer()

        @step
        def progress_counts_up_and_says_go_slack():
            m = browser.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True)
            p = m.new_page()
            p.on('pageerror', lambda e: errors.append(f'progress pageerror: {e}'))
            offline(p)
            # 谷歌慢一点、快慢不一（在浏览器里等，别堵住这边转发请求的线程）：看得见 0、1、2、3 往上跳
            p.add_init_script('''(() => { const of = window.fetch; let k = 0;
              window.fetch = async (u, o) => { if (String(u).includes('googleapis.com')) await new Promise(r => setTimeout(r, 900 + 500 * (k++ % 3))); return of(u, o); };
            })()''')
            p.goto(base)
            p.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
            p.evaluate('''() => { window.__seq = []; new MutationObserver(() => {
                const t = document.querySelector("#status").textContent; if (window.__seq.at(-1) !== t) window.__seq.push(t);
              }).observe(document.querySelector("#status"), {childList: true, subtree: true, characterData: true}); }''')
            try:
                p.set_input_files('#files', make_images(os.path.join(WORK, 'prog'), 3, start=70))
                p.fill('#key1', K1)
                p.click('#start')
                p.wait_for_function('() => /^0\\/3/.test(document.querySelector("#status").textContent)', timeout=5000)
                assert p.inner_text('#slogan') == '你可以去摸鱼了，剩下的交给老衲。', p.inner_text('#slogan')
                p.wait_for_function('() => parseFloat(document.querySelector("#bar").style.width) > 3', timeout=4000)   # 还没认完一张，条已经在爬
                p.wait_for_timeout(700)
                top = p.evaluate('() => document.querySelector(".side").getBoundingClientRect().top')
                assert 0 <= top < 844, top                                         # 手机上自己滚到任务状态
                p.screenshot(path=os.path.join(SHOTS, '7-手机干活中.png'))
                p.wait_for_function('() => document.querySelector("#status").textContent.startsWith("干完")', timeout=60000)
                seq = p.evaluate('() => window.__seq')
                ticks = [t for t in seq if re.fullmatch(r'\d+/3 张', t)]
                assert ticks == ['0/3 张', '1/3 张', '2/3 张', '3/3 张'], seq
                assert p.inner_text('#slogan') == '干完你核对就好了。', p.inner_text('#slogan')
                if p.is_visible('#dlg-view'):
                    p.click('#vclose')
                p.screenshot(path=os.path.join(SHOTS, '7-手机干完.png'))
            finally:
                m.close()

        def idb_get(pg, store, key):
            return pg.evaluate("""([s, k]) => new Promise((ok, bad) => {
                const r = indexedDB.open('luruseng-web');
                r.onsuccess = () => { const q = r.result.transaction(s).objectStore(s).get(k);
                  q.onsuccess = () => { r.result.close(); ok(q.result === undefined ? null : q.result); }; q.onerror = () => bad(q.error); };
                r.onerror = () => bad(r.error);
            })""", [store, key])

        PENDING = """() => new Promise((ok) => { const r = indexedDB.open('luruseng-web');
            r.onsuccess = () => { const q = r.result.transaction('photos').objectStore('photos').getAllKeys(IDBKeyRange.bound('pending/', 'pending/\\uffff'));
              q.onsuccess = () => { r.result.close(); ok(q.result.map(k => k.slice(8).split('|')[0])); }; q.onerror = () => ok(null); };
            r.onerror = () => ok(null); })"""

        def wait_pending(pg, n, timeout=10):  # 等存着的原照片（还没干完的）正好 n 张（wait_for_function 不等 Promise，只能自己轮询）
            end = time.time() + timeout
            while True:
                left = pg.evaluate(PENDING)
                if left is not None and len(left) == n:
                    return left
                assert time.time() < end, f'存着的照片应该是 {n} 张，现在是 {left}'
                pg.wait_for_timeout(100)

        @step
        def killed_midway_comes_back_and_finishes():
            k = browser.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True,
                                    has_touch=True, accept_downloads=True)
            # 谷歌回得有快有慢（在浏览器里等）：一批 5 张陆续认完，好在中间把页面杀掉
            k.add_init_script("""(() => { const of = window.fetch; let n = 0;
              window.fetch = async (u, o) => { if (String(u).includes('googleapis.com')) await new Promise(r => setTimeout(r, [150, 600, 1100, 1600, 2100][n++ % 5])); return of(u, o); };
            })()""")
            k.add_init_script("""document.addEventListener('DOMContentLoaded', () => { const el = document.querySelector('#slogan'); window.__slogans = [];
              new MutationObserver(() => { if (window.__slogans.at(-1) !== el.textContent) window.__slogans.push(el.textContent); })
                .observe(el, { childList: true, characterData: true, subtree: true }); })""")       # 进度条底下那句话先后说过啥

            def open_app():
                pg = k.new_page()
                pg.on('pageerror', lambda e: errors.append(f'kill pageerror: {e}'))
                pg.on('console', lambda m: errors.append(f'kill console.{m.type}: {m.text}')
                      if m.type == 'error' and 'Failed to load resource' not in m.text else None)
                offline(pg)
                pg.goto(base)
                pg.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
                return pg

            def peek(fn):                     # 页面被杀之后，浏览器里到底存下了什么（开个不跑录入僧的页面看）
                q = k.new_page()
                q.goto(base + 'icon-180.png')
                try:
                    return fn(q)
                finally:
                    q.close()

            def logs_of(pg):
                return pg.eval_on_selector_all('#log p', 'ps => ps.map(p => p.textContent)')

            try:
                p = open_app()
                p.fill('#key1', K1)
                p.set_input_files('#files', make_images(os.path.join(WORK, 'kill'), 8, start=140))
                p.click('#start')
                p.wait_for_function('() => /^2\\/8/.test(document.querySelector("#status").textContent)', timeout=30000)
                p.screenshot(path=os.path.join(SHOTS, '9-被杀之前.png'))
                p.close()                                                       # 手机把页面杀了
                work, run, left = peek(lambda q: (lambda w: (w, idb_get(q, 'runs', w['run']), q.evaluate(PENDING)))(idb_get(q, 'kv', 'work')))
                done = sorted({r['no'] for r in run['rows']})
                assert 2 <= len(done) < 8 and len(work['items']) == 8, (done, work)
                todo = {f'img_{139 + no}.jpg' for no in range(1, 9) if no not in done}
                assert todo <= set(left), (sorted(left), done)                  # 没干完的照片都还存着
                calls = dict(mock_server.STATE['calls'])

                p = open_app()                                                  # 重新打开：不点、不选，自己接着干
                p.wait_for_function('() => document.querySelector("#status").textContent.startsWith("干完")', timeout=60000)
                text = '\n'.join(logs_of(p))
                assert f'老衲诈尸接着干：还剩 {8 - len(done)} 张，接着写进 {run["xlsx"]}，编号不变' in text, text
                assert '以前干过' not in text and '没存住' not in text, text
                assert p.inner_text('#status').replace('\n', '') == '干完：成功 8 · 失败 0', p.inner_text('#status')
                assert p.inner_text('#slogan') == '干完你核对就好了。'
                said = p.evaluate('() => window.__slogans')
                assert said[0] == '页面被关了一回，老衲诈尸接着干，你接着摸鱼。' and said[-1] == '干完你核对就好了。', said
                for no in done:                                                 # 被杀之前干完的：一张都没重干
                    assert mock_server.STATE['calls'][939 + no] == calls[939 + no] == 2, (no, calls.get(939 + no))
                after = idb_get(p, 'runs', work['run'])
                assert sorted(r['no'] for r in after['rows']) == list(range(1, 9)), after['rows']   # 同一本表，1–8 各一行
                assert all(r['kind'] == 'rec' and not r.get('left') for r in after['rows'])
                assert idb_get(p, 'kv', 'work') is None                         # 这一笔了结了
                wait_pending(p, 0)                                              # 干完了：存的原照片清掉
                assert not re.match(r'\d+ 张', p.inner_text('#droptxt')), p.inner_text('#droptxt')
                p.wait_for_selector('#dlg-view[open]')
                with p.expect_download() as d:
                    p.click('#vxlsx')
                assert d.value.suggested_filename == run['xlsx']
                path = os.path.join(WORK, 'kill-' + d.value.suggested_filename)
                d.value.save_as(path)
                ws = load_workbook(path).worksheets[0]
                assert [ws.cell(row=r, column=1).value for r in range(4, 12)] == list(range(1, 9))
                assert [int(ws.cell(row=r, column=6).value) for r in range(4, 12)] == [940 + i for i in range(8)], \
                    [ws.cell(row=r, column=6).value for r in range(4, 12)]          # 照片和编号对得上
                p.click('#vclose')
                p.screenshot(path=os.path.join(SHOTS, '9-诈尸干完.png'))

                # 点了开始、照片还没存完页面就被杀（在 iPhone 上：切走马上被冻住，后来被杀）：存住的自己干完，没存住的点名
                p.set_input_files('#files', make_images(os.path.join(WORK, 'half'), 1, start=160))
                kept = wait_pending(p, 1)
                ids = p.evaluate("""() => new Promise((ok) => { const r = indexedDB.open('luruseng-web');
                    r.onsuccess = () => { const q = r.result.transaction('photos').objectStore('photos').getAllKeys(IDBKeyRange.bound('pending/', 'pending/\\uffff'));
                      q.onsuccess = () => { r.result.close(); ok(q.result.map(k => k.slice(8))); }; }; })""")
                p.evaluate("""want => new Promise((ok) => { const r = indexedDB.open('luruseng-web');
                    r.onsuccess = () => { const t = r.result.transaction('kv', 'readwrite');
                      t.objectStore('kv').put({ at: Date.now(), auto: false, stuck: 0, want }, 'work'); t.oncomplete = () => { r.result.close(); ok(); }; }; })""",
                           ids + ['ghost.jpg|12345|1700000000000'])
                p.close()
                p = open_app()
                p.wait_for_function('() => document.querySelector("#status").textContent.startsWith("干完")', timeout=60000)
                text = '\n'.join(logs_of(p))
                assert '▶ 上回刚开工页面就被关了，老衲诈尸重新开工' in text and '⚠ 这 1 张还没存进浏览器页面就被关了，得重选：ghost.jpg' in text, text
                assert kept == ['img_160.jpg'] and p.inner_text('#status').replace('\n', '') == '干完：成功 1 · 失败 0', p.inner_text('#status')
                assert idb_get(p, 'kv', 'work') is None
                wait_pending(p, 0)
                if p.is_visible('#dlg-view'):
                    p.click('#vclose')

                p.set_input_files('#files', make_images(os.path.join(WORK, 'keep'), 2, start=150))   # 选了没开工就被关
                assert sorted(wait_pending(p, 2)) == ['img_150.jpg', 'img_151.jpg']
                p.close()
                p = open_app()
                p.wait_for_function('() => [...document.querySelectorAll("#log p")].some(p => p.textContent.includes("上回留下的 2 张照片还在列表里"))')
                assert p.inner_text('#droptxt').startswith('2 张：img_150.jpg'), p.inner_text('#droptxt')
                p.click('#clear')
                wait_pending(p, 0)
                p.close()
                p = open_app()
                assert not re.match(r'\d+ 张', p.inner_text('#droptxt'))
                assert not any('上回' in m for m in logs_of(p)), logs_of(p)
            finally:
                k.close()

        @step
        def phone_sends_and_pc_starts_by_itself():
            if not relay_url:
                print('    （没有中转站代码，手机传图不测）')
                return
            init = f"window.LURUSENG_RELAY = '{relay_url}';"
            pc = browser.new_context(viewport={'width': 1280, 'height': 860}, accept_downloads=True, locale='zh-CN')
            ph = browser.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True)
            for c in (pc, ph):
                c.add_init_script(init)

            def opened(ctx, url, wait_log=True):
                pg = ctx.new_page()
                pg.on('pageerror', lambda e: errors.append(f'phone-test pageerror: {e}'))
                pg.on('console', lambda m: errors.append(f'phone-test console.{m.type}: {m.text}')
                      if m.type == 'error' and 'Failed to load resource' not in m.text else None)
                offline(pg)
                pg.goto(url)
                if wait_log:
                    pg.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
                return pg

            def logs_of(pg):
                return '\n'.join(pg.eval_on_selector_all('#log p', 'ps => ps.map(p => p.textContent)'))

            try:
                a = opened(pc, base)
                assert a.is_visible('#phone') and '手机扫个码' in a.inner_text('#phinttxt')
                a.fill('#key1', K1)
                a.click('#phone')
                a.wait_for_selector('#dlg-phone[open]')
                a.wait_for_function('() => !!document.querySelector("#qr svg")')
                link = a.get_attribute('#qr', 'data-link')
                assert '/web/send.html#r=' in link and '&k=' in link, link
                room = link.split('#r=')[1].split('&')[0]
                a.wait_for_function('() => document.querySelector("#pstate").textContent.includes("已配对")', timeout=15000)
                a.screenshot(path=os.path.join(SHOTS, '10-手机传图二维码.png'))
                a.click('#dlg-phone [data-close]')

                b = opened(ph, link, wait_log=False)                            # 手机扫码打开的那一页
                b.wait_for_function('() => document.querySelector("#pcstate").textContent.includes("电脑在线")', timeout=15000)
                b.set_input_files('#files', make_images(os.path.join(WORK, 'ph'), 3, start=170))
                b.wait_for_function('() => document.querySelector("#st").textContent.includes("电脑收齐 3 张")', timeout=30000)
                a.wait_for_function('() => document.querySelector("#status").textContent.startsWith("干完")', timeout=60000)
                text = logs_of(a)
                assert '📱 手机传来 3 张，收齐了，开工' in text, text
                assert a.inner_text('#status').replace('\n', '') == '干完：成功 3 · 失败 0', a.inner_text('#status')
                b.wait_for_function('() => document.querySelector("#say").textContent.includes("干完了，成功 3 · 失败 0")', timeout=15000)
                b.screenshot(path=os.path.join(SHOTS, '10-手机传完.png'))
                run_id = a.evaluate("""() => new Promise((ok) => { const r = indexedDB.open('luruseng-web');
                    r.onsuccess = () => { const q = r.result.transaction('kv').objectStore('kv').get('cfg'); q.onsuccess = () => ok(q.result.lastRun); }; })""")
                rows = a.evaluate("""id => new Promise((ok) => { const r = indexedDB.open('luruseng-web');
                    r.onsuccess = () => { const q = r.result.transaction('runs').objectStore('runs').get(id); q.onsuccess = () => ok(q.result.rows.map(x => [x.no, x.values.opr_no])); }; })""", run_id)
                assert [[n, str(v)] for n, v in rows] == [[1, '970'], [2, '971'], [3, '972']], rows   # 按拍摄顺序编号，一张不少
                left = a.evaluate(f"fetch('{relay_url}/r/{room}/items').then(r => r.json())")
                assert left == [], left                                                # 取走就删
                if a.is_visible('#dlg-view'):
                    a.click('#vclose')

                a.close()                                                              # 电脑关了：照片先存在中转站
                b.wait_for_function('() => document.querySelector("#pcstate").textContent.includes("电脑没开")', timeout=15000)
                b.set_input_files('#files', make_images(os.path.join(WORK, 'ph2'), 2, start=175))
                b.wait_for_function('() => document.querySelector("#st").textContent.includes("送出 2/2")', timeout=30000)
                assert '照片在中转站等着' in b.inner_text('#say'), b.inner_text('#say')
                a = opened(pc, base)                                                   # 电脑再打开：自己取、自己干
                a.wait_for_function('() => document.querySelector("#status").textContent.startsWith("干完")', timeout=60000)
                assert '📱 手机传来 2 张，收齐了，开工' in logs_of(a), logs_of(a)
                b.wait_for_function('() => document.querySelector("#st").textContent.includes("电脑收齐 2 张")', timeout=15000)
                assert a.evaluate(f"fetch('{relay_url}/r/{room}/items').then(r => r.json())") == []
            finally:
                pc.close()
                ph.close()

        @step
        def phone_layout():
            m = browser.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True)
            p = m.new_page()
            p.on('pageerror', lambda e: errors.append(f'phone pageerror: {e}'))
            offline(p)
            p.goto(base)
            p.wait_for_function('() => document.querySelectorAll("#log p").length >= 1')
            sw = p.evaluate('() => [document.documentElement.scrollWidth, window.innerWidth]')
            assert sw[0] <= sw[1], sw                                             # 不横着滚
            p.screenshot(path=os.path.join(SHOTS, '5-手机.png'), full_page=True)
            for btn, dlg in (('#gear', '#dlg-set'), ('#names', '#dlg-names'), ('#feed', '#dlg-feed'), ('#newcat', '#dlg-new')):
                p.click(btn)
                p.wait_for_selector(f'{dlg}[open]')
                box = p.evaluate('d => { const r = document.querySelector(d + " .dcard").getBoundingClientRect(); return [r.left, r.right, innerWidth]; }', dlg)
                assert 0 <= box[0] and box[1] <= box[2], (dlg, box)              # 对话框整个在屏幕里，左边不被切掉
                if dlg == '#dlg-set':
                    p.screenshot(path=os.path.join(SHOTS, '5-手机设置.png'))
                p.click(f'{dlg} [data-close]')
                p.wait_for_selector(f'{dlg}:not([open])', state='attached')
            p.set_input_files('#files', imgs[:3])
            p.fill('#key1', K1)
            p.click('#start')
            p.wait_for_selector('#dlg-view[open]', timeout=60000)
            p.wait_for_function('() => { const i = document.querySelector("#vphoto img"); return i && i.naturalWidth > 0; }')
            sw = p.evaluate('() => [document.documentElement.scrollWidth, window.innerWidth]')
            assert sw[0] <= sw[1], sw
            p.screenshot(path=os.path.join(SHOTS, '6-手机核对.png'))
            hs = p.evaluate("""() => Object.fromEntries([...document.querySelectorAll('#vpanel dt')]
                .map(d => [d.textContent, Math.round(d.getBoundingClientRect().height)]))""")
            assert hs['TOTAL SCREENED'] == hs['MACHINE'], hs                  # 手机上 TOTAL SCREENED 也是一行
            p.click('#vpanel dd[data-f="machine"]')                            # 手机上点格子改
            p.wait_for_selector('#vpanel .cell-ed input')
            sw = p.evaluate('() => [document.documentElement.scrollWidth, window.innerWidth]')
            assert sw[0] <= sw[1], sw
            p.eval_on_selector('#vpanel .cell-ed', 'e => e.scrollIntoView({block: "center"})')
            p.screenshot(path=os.path.join(SHOTS, '6-手机改格子.png'))
            m.close()

        for fn in (opens_clean, no_key_no_start, twelve_photos_match_desktop_book, viewer_pages_rotates_copies_zips,
                   same_photos_again_are_skipped, result_button_downloads_latest, feed_database_xlsx_xls_csv_html,
                   database_fixes_confusable_letters, edit_cells_in_viewer, btmp_copies_total, names_dialog, own_category_matches_desktop_book,
                   bad_key_stops_and_keeps_photos, stop_button_marks_unfinished, photos_of_all_kinds,
                   title_on_the_right_turns_upright_number_in_top_right, settings_backup_restore_and_forget_keys, delete_a_batch, progress_counts_up_and_says_go_slack,
                   killed_midway_comes_back_and_finishes, phone_sends_and_pc_starts_by_itself, phone_layout):
            if os.environ.get('E2E_ONLY') and fn.__name__ not in os.environ['E2E_ONLY'].split(','):
                continue                      # 只跑某几项：E2E_ONLY=项名,项名
            try:
                fn()
            except Exception:
                print(f'  ✗ {fn.__name__}', flush=True)
                print(last_logs(12))
                page.screenshot(path=os.path.join(SHOTS, f'失败-{fn.__name__}.png'))
                raise
        page.screenshot(path=os.path.join(SHOTS, '8-最后.png'))
        browser.close()
    if relay_proc:
        os.killpg(relay_proc.pid, signal.SIGTERM)
    bad = [e for e in errors if 'favicon' not in e]
    assert not bad, '\n'.join(bad)
    print(f'全过：{len(passed)} 项 · 截图在 {SHOTS}')


if __name__ == '__main__':
    main()
